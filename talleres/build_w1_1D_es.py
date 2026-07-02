# -*- coding: utf-8 -*-
"""Taller 1 — Cinemática en una dimensión (1D). Resuelto en ESPAÑOL."""
import docbuild as D
from docbuild import (new_doc, cover, exercise_heading, statement, subhead,
                      para, bullet, add_eq, answer, answer_text, image, _shade)
from docbuild import r, i, sup, sub, subsup, frac, sqrt, paren, brack
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

NAME = "EDWARD ANDRÉS RODRÍGUEZ MAZUERA"
CODE = "2535250-2713"

def data_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    hdr = t.rows[0].cells
    for c, h in zip(hdr, headers):
        _shade(c, '1F3A5F'); p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p.add_run(h); rr.bold = True; rr.font.color.rgb = RGBColor(0xFF,0xFF,0xFF); rr.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(v).font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

doc = new_doc()
cover(doc, "Cinemática en una Dimensión (1D)",
      "Capítulo 2 — Ejercicios (Clase 03)", NAME, CODE, lang='es')

para(doc, "Método de siete pasos: (1) pensar, (2) esquematizar, (3) investigar, "
          "(4) simplificar, (5) calcular, (6) redondear, (7) analizar. Se trabaja "
          "en el SI y se usa g = 9.8 m/s².",
     size=10, italic=True, space_after=8)

# ---------------------------------------------------------------- Ej 1
exercise_heading(doc, "Ejercicio 1")
statement(doc, "La gráfica velocidad–tiempo de un objeto que parte del reposo (en "
               "el origen) y se mueve a lo largo del eje x se muestra en la figura. "
               "(a) Determinar la aceleración media en los intervalos (0, 5.0) s, "
               "(5.0, 15.0) s y (15.0, 20.0) s. (b) Trazar la curva aceleración–tiempo. "
               "(c) Trazar la curva posición–tiempo. (d) En t = 10.0 s, ¿cuál fue el "
               "desplazamiento?")
image(doc, "img/w1e1_vt.png", 9.5, "Figura. Gráfica velocidad–tiempo dada (v pasa de −8 m/s a +8 m/s).")
subhead(doc, "(a)  Aceleración media  a = Δv/Δt")
add_eq(doc, sub(i('a'),r('1')), r('='), frac(r('−8 −'), r('')) if False else frac(r('Δv'),r('Δt')),
            r('='), frac(r('0'), r('5.0')), r('='), r('0 m/s²'), r('     (v constante = −8 m/s)'))
add_eq(doc, sub(i('a'),r('2')), r('='), frac(r('8 − (−8)'), r('15.0 − 5.0')), r('='),
            frac(r('16'), r('10')), r('='), r('1.6 m/s²'))
add_eq(doc, sub(i('a'),r('3')), r('='), frac(r('8 − 8'), r('20.0 − 15.0')), r('='), r('0 m/s²'))
data_table(doc, ["Intervalo (s)", "(0, 5.0)", "(5.0, 15.0)", "(15.0, 20.0)"],
                [["a (m/s²)", "0", "1.6", "0"]])
subhead(doc, "(b)  Curva aceleración–tiempo")
image(doc, "img/w1e1_at.png", 9.5, "Figura. Aceleración–tiempo: constante a tramos.")
subhead(doc, "(c)  Curva posición–tiempo")
para(doc, "Se integra la velocidad en cada tramo. La función v(t) es:")
add_eq(doc, sub(i('v'),i('x')), r('='), r('−8'), r('  (0<t<5) ;   '),
            r('1.6'), i('t'), r('−16'), r('  (5<t<15) ;   '), r('8'), r('  (15<t<20)'))
para(doc, "Integrando y ajustando constantes por continuidad (x = 0 en t = 0):")
add_eq(doc, i('x'), r('='), r('−8'), i('t'), r('  (0<t<5) ;   '),
            r('0.8'), sup(i('t'),r('2')), r('−16'), i('t'), r('+20'), r('  (5<t<15) ;   '),
            r('8'), i('t'), r('−160'), r('  (15<t<20)'))
image(doc, "img/w1e1_xt.png", 9.5, "Figura. Posición–tiempo resultante.")
subhead(doc, "(d)  Desplazamiento en t = 10.0 s")
add_eq(doc, i('x'), paren(r('10')), r('='), r('0.8'), sup(r('(10)'),r('2')), r('−16'), r('(10)'), r('+20'),
            r('='), r('−60 m'))
answer(doc, i('x'), paren(r('10 s')), r('='), r('−60 m'), label='Respuesta (d):')

# ---------------------------------------------------------------- Ej 2
exercise_heading(doc, "Ejercicio 2")
statement(doc, "Una partícula parte del reposo y acelera como se muestra en la "
               "figura. Determine (a) la velocidad de la partícula en t = 10.0 s y en "
               "t = 20.0 s; (b) la distancia recorrida en los primeros 20.0 s.")
image(doc, "img/w1e2_at.png", 9.5, "Figura. Gráfica aceleración–tiempo dada.")
subhead(doc, "Función aceleración")
add_eq(doc, sub(i('a'),i('x')), r('='), r('2'), r('  (0<t<10) ;   '),
            r('0'), r('  (10<t<15) ;   '), r('−3'), r('  (15<t<20)'))
subhead(doc, "(a)  Velocidad — integración de a(t)")
add_eq(doc, sub(i('v'),i('x')), r('='), r('2'), i('t'), r('  (0<t<10) ;   '),
            r('20'), r('  (10<t<15) ;   '), r('−3'), i('t'), r('+65'), r('  (15<t<20)'))
para(doc, "Constantes: k₁ = 0 (parte del reposo); k₂ = 20 (continuidad en t = 10); "
          "k₃ = 65 (continuidad en t = 15). Por tanto:")
add_eq(doc, i('v')+paren(r('10')), r('='), r('2(10)'), r('='), r('20 m/s'), r('   ,   '),
            i('v')+paren(r('20')), r('='), r('−3(20)+65'), r('='), r('5 m/s'))
answer_text(doc, "v(10 s) = 20 m/s ;   v(20 s) = 5 m/s.", label='Respuesta (a):')
image(doc, "img/w1e2_vt.png", 9.5, "Figura. Gráfica velocidad–tiempo obtenida.")
subhead(doc, "(b)  Distancia — área bajo la curva v(t)")
add_eq(doc, i('d'), r('='), frac(r('1'),r('2')), r('(10)(20)'), r('+'), r('(5)(20)'), r('+'),
            frac(r('1'),r('2')), r('(20+5)(5)'))
add_eq(doc, i('d'), r('='), r('100 + 100 + 62.5'), r('='), r('262.5 m'))
image(doc, "img/w1e2_xt.png", 9.5, "Figura. Gráfica posición–tiempo obtenida.")
answer(doc, i('d'), r('='), r('262.5 m'), label='Respuesta (b):')

# ---------------------------------------------------------------- Ej 3
exercise_heading(doc, "Ejercicio 3")
statement(doc, "Un tren viaja entre dos estaciones separadas 1.00 km y nunca alcanza "
               "su velocidad de crucero máxima. En horas pico el ingeniero minimiza el "
               "tiempo Δt acelerando durante Δt₁ a razón de a₁ = 0.100 m/s² e "
               "inmediatamente desacelerando a a₂ = −0.500 m/s² durante Δt₂. Encontrar "
               "el tiempo mínimo de viaje Δt y el intervalo Δt₁.")
subhead(doc, "Datos y relaciones")
bullet(doc, "x = 1.00 km = 1000 m,  a₁ = 0.100 m/s²,  a₂ = −0.500 m/s²")
para(doc, "El tren parte del reposo y se detiene al final (v₂ = 0). Como "
          "v₁ = a₁t₁ = −a₂t₂, se tiene t₂ = 0.2 t₁. La distancia total es:")
add_eq(doc, i('x'), r('='), sub(i('x'),r('1')), r('+'), sub(i('x'),r('2')), r('='),
            frac(r('1'),r('2')), sub(i('a'),r('1')), subsup(i('t'),r('1'),r('2')), r('+'),
            sub(i('v'),r('1')), sub(i('t'),r('2')), r('+'), frac(r('1'),r('2')), sub(i('a'),r('2')),
            subsup(i('t'),r('2'),r('2')))
para(doc, "Sustituyendo t₂ = 0.2 t₁ y v₁ = 0.1 t₁ se llega a 0.06 t₁² = 1000, luego:")
add_eq(doc, sub(i('t'),r('1')), r('='), sqrt(frac(r('1000'), r('0.06'))), r('='),
            sqrt(frac(r('20000'), r('1.20'))), r('='), r('129 s'))
add_eq(doc, sub(i('t'),r('2')), r('='), frac(sub(i('a'),r('1'))+sub(i('t'),r('1')), r('−')+sub(i('a'),r('2'))),
            r('='), frac(r('(0.100)(129)'), r('0.500')), r('='), r('25.8 s'))
add_eq(doc, r('Δ')+i('t'), r('='), sub(i('t'),r('1')), r('+'), sub(i('t'),r('2')), r('='),
            r('129 + 25.8'), r('≈'), r('155 s'))
answer(doc, r('Δ')+i('t'), r('≈'), r('155 s'), r('  ,  '), r('Δ')+sub(i('t'),r('1')), r('='),
       r('129 s'), label='Respuesta:')

# ---------------------------------------------------------------- Ej 4
exercise_heading(doc, "Ejercicio 4")
statement(doc, "María compra un carro deportivo que acelera a 4.90 m/s². Decide "
               "probarlo contra José. Ambos parten del reposo, pero José arranca 1.00 s "
               "antes que María. Si José acelera a 3.50 m/s² y María mantiene su "
               "aceleración de 4.90 m/s², hallar (a) el tiempo en que María alcanza a "
               "José, (b) la distancia que ella recorre antes de alcanzarlo y (c) las "
               "velocidades de ambos carros en ese instante.")
subhead(doc, "Datos")
bullet(doc, "a_M = 4.90 m/s²,  a_J = 3.50 m/s².  Condición: x_M = x_J con t_J = t_M + 1.00")
subhead(doc, "(a)  Tiempo en que María alcanza a José")
add_eq(doc, frac(r('1'),r('2')), r('(3.50)'), sup(paren(sub(i('t'),i('M'))+r('+1.00')),r('2')), r('='),
            frac(r('1'),r('2')), r('(4.90)'), subsup(i('t'),i('M'),r('2')))
add_eq(doc, r('1.40'), subsup(i('t'),i('M'),r('2')), r('−'), r('7.00'), sub(i('t'),i('M')), r('−'),
            r('3.50'), r('='), r('0'), r('  ⇒  '), sub(i('t'),i('M')), r('='), r('5.46 s'))
answer_text(doc, "t_M ≈ 5.5 s.", label='Respuesta (a):')
subhead(doc, "(b)  Distancia recorrida por María")
add_eq(doc, sub(i('x'),i('M')), r('='), frac(r('1'),r('2')), r('(4.90)'), sup(r('(5.46)'),r('2')),
            r('='), r('73.0 m'))
answer_text(doc, "x_M ≈ 73.0 m.", label='Respuesta (b):')
subhead(doc, "(c)  Velocidades")
add_eq(doc, sub(i('v'),i('M')), r('='), r('(4.90)(5.46)'), r('='), r('26.7 m/s'), r('   ,   '),
            sub(i('v'),i('J')), r('='), r('(3.50)(6.46)'), r('='), r('22.6 m/s'))
answer(doc, sub(i('v'),i('M')), r('≈'), r('26.7 m/s'), r('  ,  '), sub(i('v'),i('J')), r('≈'),
       r('22.6 m/s'), label='Respuesta (c):')

# ---------------------------------------------------------------- Ej 5
exercise_heading(doc, "Ejercicio 5")
statement(doc, "Una piedra se deja caer desde el reposo en un pozo. El sonido del "
               "chapoteo se escucha 2.40 s después de soltar la roca. (a) ¿A qué "
               "distancia por debajo de la parte superior del pozo está la superficie "
               "del agua? La velocidad del sonido en el aire es 336 m/s. (b) ¿Qué pasa "
               "si no se tiene en cuenta la rapidez del sonido? (c) ¿Qué porcentaje de "
               "error se comete?")
subhead(doc, "Planteamiento")
para(doc, "El tiempo total se reparte entre la caída libre (t₁) y el regreso del "
          "sonido (t₂):  t₁ + t₂ = 2.40 s, con d = ½g t₁² y d = 336 t₂.")
subhead(doc, "(a)  Distancia real")
add_eq(doc, r('4.90'), sup(paren(r('2.40 −')+sub(i('t'),r('2'))),r('2')), r('='), r('336'), sub(i('t'),r('2')),
            r('  ⇒  '), r('4.90'), subsup(i('t'),r('2'),r('2')), r('−'), r('359.5'), sub(i('t'),r('2')), r('+'),
            r('28.22'), r('='), r('0'))
add_eq(doc, sub(i('t'),r('2')), r('='), r('0.0786 s'), r('  ⇒  '), i('d'), r('='), r('336'),
            paren(r('0.0786')), r('='), r('26.4 m'))
answer(doc, i('d'), r('≈'), r('26.4 m'), label='Respuesta (a):')
subhead(doc, "(b)  Ignorando la rapidez del sonido")
add_eq(doc, i('d'), r('='), frac(r('1'),r('2')), r('(9.80)'), sup(r('(2.40)'),r('2')), r('='), r('28.2 m'))
answer_text(doc, "d ≈ 28.2 m.", label='Respuesta (b):')
subhead(doc, "(c)  Porcentaje de error")
add_eq(doc, r('error'), r('='), frac(r('28.2 − 26.4'), r('26.4')), r('×'), r('100%'), r('='), r('6.8%'))
answer(doc, r('error'), r('≈'), r('6.8%'), label='Respuesta (c):')

# ---------------------------------------------------------------- Ej 6
exercise_heading(doc, "Ejercicio 6")
statement(doc, "Un objeto en caída libre tarda 1.5 s en recorrer los últimos 30.0 m "
               "antes de golpear el suelo. ¿Desde qué altura sobre el suelo cayó?")
subhead(doc, "Solución")
para(doc, "Sea v la rapidez al inicio de ese último tramo. Durante los últimos 30 m:")
add_eq(doc, r('30'), r('='), i('v'), paren(r('1.5')), r('+'), frac(r('1'),r('2')), r('(9.8)'),
            sup(r('(1.5)'),r('2')), r('  ⇒  '), i('v'), r('='), r('12.65 m/s'))
para(doc, "Esa rapidez se alcanzó tras caer una altura h₁ desde el reposo:")
add_eq(doc, sub(i('h'),r('1')), r('='), frac(sup(i('v'),r('2')), r('2')+i('g')), r('='),
            frac(sup(r('(12.65)'),r('2')), r('2(9.8)')), r('='), r('8.16 m'))
add_eq(doc, i('H'), r('='), sub(i('h'),r('1')), r('+'), r('30.0'), r('='), r('38.2 m'))
answer(doc, i('H'), r('≈'), r('38.2 m'), label='Respuesta:')

# ---------------------------------------------------------------- Ej 7
exercise_heading(doc, "Ejercicio 7")
statement(doc, "Un estudiante lanza un juego de llaves verticalmente hacia arriba a "
               "un compañero que está en una ventana 4.00 m más arriba. La mano del "
               "compañero atrapa las llaves 1.50 s después. (a) ¿Con qué velocidad "
               "inicial se lanzaron las llaves? (b) ¿Cuál era la velocidad de las "
               "llaves justo antes de ser atrapadas?")
subhead(doc, "(a)  Velocidad inicial")
add_eq(doc, i('y'), r('='), sub(i('v'),r('0')), i('t'), r('−'), frac(r('1'),r('2')), i('g'), sup(i('t'),r('2')),
            r('  ⇒  '), r('4.00'), r('='), sub(i('v'),r('0')), paren(r('1.50')), r('−'), r('11.025'))
add_eq(doc, sub(i('v'),r('0')), r('='), frac(r('4.00 + 11.025'), r('1.50')), r('='), r('10.0 m/s'))
answer(doc, sub(i('v'),r('0')), r('='), r('10.0 m/s'), label='Respuesta (a):')
subhead(doc, "(b)  Velocidad justo antes de atraparlas")
add_eq(doc, i('v'), r('='), sub(i('v'),r('0')), r('−'), i('g'), i('t'), r('='), r('10.0 − (9.8)(1.50)'),
            r('='), r('−4.7 m/s'))
answer(doc, i('v'), r('≈'), r('−4.7 m/s'), r('  (4.7 m/s hacia abajo)'), label='Respuesta (b):')

# ---------------------------------------------------------------- Ej 8
exercise_heading(doc, "Ejercicio 8")
statement(doc, "Dos carros, A y B, se mueven en la misma dirección. En t = 0 sus "
               "velocidades son 1 m/s y 3 m/s, y sus aceleraciones son 2 m/s² y 1 m/s², "
               "respectivamente. Si el carro A está 1.5 m adelante de B en t = 0, "
               "calcule cuándo y dónde estarán lado a lado.")
subhead(doc, "Ecuaciones de posición")
add_eq(doc, sub(i('x'),i('A')), r('='), r('1.5'), r('+'), i('t'), r('+'), sup(i('t'),r('2')),
            r('   ,   '), sub(i('x'),i('B')), r('='), r('3'), i('t'), r('+'), r('0.5'), sup(i('t'),r('2')))
para(doc, "Están lado a lado cuando x_A = x_B:")
add_eq(doc, r('1.5 +'), i('t'), r('+'), sup(i('t'),r('2')), r('='), r('3'), i('t'), r('+'), r('0.5'), sup(i('t'),r('2')),
            r('  ⇒  '), r('0.5'), sup(i('t'),r('2')), r('−'), r('2'), i('t'), r('+'), r('1.5'), r('='), r('0'))
add_eq(doc, sup(i('t'),r('2')), r('−'), r('4'), i('t'), r('+'), r('3'), r('='), r('0'), r('  ⇒  '),
            paren(i('t')+r('−1')), paren(i('t')+r('−3')), r('='), r('0'))
para(doc, "Las soluciones son t = 1.0 s y t = 3.0 s. Las posiciones correspondientes son:")
add_eq(doc, i('x')+paren(r('1 s')), r('='), r('3.5 m'), r('   ,   '), i('x')+paren(r('3 s')), r('='), r('13.5 m'))
answer(doc, i('t'), r('='), r('1.0 s'), r(' (x = 3.5 m)'), r('   y   '), i('t'), r('='), r('3.0 s'),
       r(' (x = 13.5 m)'), label='Respuesta:')

# ---------------------------------------------------------------- Ej 9
exercise_heading(doc, "Ejercicio 9")
statement(doc, "La pelota A se deja caer desde el reposo desde una altura h sobre el "
               "suelo. La pelota B se lanza verticalmente hacia arriba desde el suelo en "
               "el instante en que se suelta A. ¿Cuál es la rapidez de B si las dos "
               "pelotas deben encontrarse a una altura h/2 sobre el suelo?")
subhead(doc, "Solución")
para(doc, "Pelota A (caída libre desde h):  y_A = h − ½g t². En y = h/2:")
add_eq(doc, frac(i('h'),r('2')), r('='), i('h'), r('−'), frac(r('1'),r('2')), i('g'), sup(i('t'),r('2')),
            r('  ⇒  '), i('t'), r('='), sqrt(frac(i('h'), i('g'))))
para(doc, "Pelota B (lanzada hacia arriba):  y_B = v_B t − ½g t². En el encuentro y_B = h/2 "
          "y como ½g t² = h/2:")
add_eq(doc, frac(i('h'),r('2')), r('='), sub(i('v'),i('B')), i('t'), r('−'), frac(i('h'),r('2')),
            r('  ⇒  '), sub(i('v'),i('B')), i('t'), r('='), i('h'))
add_eq(doc, sub(i('v'),i('B')), r('='), frac(i('h'), i('t')), r('='), frac(i('h'), sqrt(frac(i('h'),i('g')))),
            r('='), sqrt(i('g')+i('h')))
answer(doc, sub(i('v'),i('B')), r('='), sqrt(i('g')+i('h')), label='Respuesta:')

out = "/home/user/AgenteSAE/talleres/Taller_Cinematica_1D_Espanol.docx"
doc.save(out)
print("Saved", out, "-- paragraphs:", len(doc.paragraphs))
