# -*- coding: utf-8 -*-
"""Helpers to build professional physics-workshop Word documents with real,
editable Word (OMML) equations, styled headings, data blocks and answer boxes."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import parse_xml, OxmlElement
from xml.sax.saxutils import escape

MNS = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'

# ---- Colour palette -------------------------------------------------------
NAVY   = RGBColor(0x1F, 0x3A, 0x5F)
BLUE   = RGBColor(0x2E, 0x5C, 0x8A)
ACCENT = RGBColor(0xC0, 0x39, 0x2B)  # deep red accent
GREEN  = RGBColor(0x1E, 0x6B, 0x3A)
GREY   = RGBColor(0x55, 0x55, 0x55)

# ==========================================================================
#  OMML equation builders  (fragments are raw XML strings inside <m:oMath>)
# ==========================================================================
def r(text, nor=False):
    """A math run. nor=True -> upright (for sin, cos, numbers, units)."""
    rpr = '<m:rPr><m:sty m:val="p"/></m:rPr>' if nor else ''
    return f'<m:r>{rpr}<m:t xml:space="preserve">{escape(str(text))}</m:t></m:r>'

def i(text):
    """Italic math variable run (default math style)."""
    return r(text, nor=False)

def sup(base, exp):
    return f'<m:sSup><m:e>{base}</m:e><m:sup>{exp}</m:sup></m:sSup>'

def sub(base, s):
    return f'<m:sSub><m:e>{base}</m:e><m:sub>{s}</m:sub></m:sSub>'

def subsup(base, s, e):
    return f'<m:sSubSup><m:e>{base}</m:e><m:sub>{s}</m:sub><m:sup>{e}</m:sup></m:sSubSup>'

def frac(num, den):
    return f'<m:f><m:num>{num}</m:num><m:den>{den}</m:den></m:f>'

def sqrt(x):
    return ('<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr>'
            f'<m:deg/><m:e>{x}</m:e></m:rad>')

def paren(x):
    return ('<m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr>'
            f'<m:e>{x}</m:e></m:d>')

def brack(x):
    return ('<m:d><m:dPr><m:begChr m:val="["/><m:endChr m:val="]"/></m:dPr>'
            f'<m:e>{x}</m:e></m:d>')

def add_eq(doc, *frags, align='center'):
    """Add a display equation paragraph built from OMML fragments."""
    body = ''.join(frags)
    xml = (f'<m:oMathPara {MNS}><m:oMath>{body}</m:oMath></m:oMathPara>')
    p = doc.add_paragraph()
    p.alignment = {'center': WD_ALIGN_PARAGRAPH.CENTER,
                   'left': WD_ALIGN_PARAGRAPH.LEFT}[align]
    p._p.append(parse_xml(xml))
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    return p

# ==========================================================================
#  Document-level styling helpers
# ==========================================================================
def new_doc():
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.4)
    return doc

def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def cover(doc, title, subtitle, name, code, lang='es'):
    L = {'es': dict(uni='Universidad del Valle',
                    course='Física y Laboratorio I',
                    st='Estudiante', cd='Código', dt='Fecha'),
         'en': dict(uni='Universidad del Valle',
                    course='Physics and Laboratory I',
                    st='Student', cd='Student code', dt='Date')}[lang]
    # University
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(L['uni'].upper()); run.bold = True; run.font.size = Pt(16)
    run.font.color.rgb = NAVY
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(L['course']); r2.font.size = Pt(12); r2.font.color.rgb = GREY
    p2.paragraph_format.space_after = Pt(14)

    # Title band (table used as a coloured bar)
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.columns[0].width = Cm(16)
    c = t.cell(0, 0); _shade(c, '1F3A5F')
    c.width = Cm(16)
    cp = c.paragraphs[0]; cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cp.add_run(title); cr.bold = True; cr.font.size = Pt(20)
    cr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    cp.paragraph_format.space_before = Pt(6); cp.paragraph_format.space_after = Pt(6)
    if subtitle:
        sp = c.add_paragraph(); sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sr = sp.add_run(subtitle); sr.font.size = Pt(12)
        sr.font.color.rgb = RGBColor(0xD6, 0xE4, 0xF0)
        sp.paragraph_format.space_after = Pt(6)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Info table
    info = doc.add_table(rows=3, cols=2); info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.style = 'Table Grid'
    rows = [(L['st'], name), (L['cd'], code), (L['dt'], '2 de julio de 2026' if lang=='es' else 'July 2, 2026')]
    for k, (lbl, val) in enumerate(rows):
        a, b = info.rows[k].cells
        _shade(a, 'E8EEF4')
        ap = a.paragraphs[0]; ar = ap.add_run(lbl); ar.bold = True; ar.font.color.rgb = NAVY
        bp = b.paragraphs[0]; br = bp.add_run(val)
    info.columns[0].width = Cm(4.5); info.columns[1].width = Cm(11.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def rule(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '1F3A5F')
    pbdr.append(bottom); pPr.append(pbdr)

def exercise_heading(doc, label):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)
    run = p.add_run(label); run.bold = True; run.font.size = Pt(14)
    run.font.color.rgb = ACCENT
    # bottom border
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '10')
    bottom.set(qn('w:space'), '2'); bottom.set(qn('w:color'), 'C0392B')
    pbdr.append(bottom); pPr.append(pbdr)

def statement(doc, text):
    """Problem statement in an indented, shaded quote box."""
    t = doc.add_table(rows=1, cols=1)
    t.columns[0].width = Cm(16)
    c = t.cell(0, 0); _shade(c, 'F2F5F8'); c.width = Cm(16)
    p = c.paragraphs[0]
    run = p.add_run(text); run.italic = True; run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    # left accent border
    tcPr = c._tc.get_or_add_tcPr(); borders = OxmlElement('w:tcBorders')
    left = OxmlElement('w:left'); left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '18'); left.set(qn('w:space'), '0'); left.set(qn('w:color'), '2E5C8A')
    borders.append(left); tcPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def subhead(doc, text, color=BLUE):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text); run.bold = True; run.font.size = Pt(11.5)
    run.font.color.rgb = color
    return p

def para(doc, text, size=11, italic=False, space_after=4):
    p = doc.add_paragraph()
    run = p.add_run(text); run.font.size = Pt(size); run.italic = italic
    p.paragraph_format.space_after = Pt(space_after)
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(text)
    p.paragraph_format.space_after = Pt(2)
    return p

def answer(doc, *frags, label='Respuesta:'):
    """Highlighted final-answer box containing an equation."""
    t = doc.add_table(rows=1, cols=1); t.columns[0].width = Cm(16)
    c = t.cell(0, 0); _shade(c, 'E7F3EC'); c.width = Cm(16)
    p = c.paragraphs[0]
    lr = p.add_run(label + '  '); lr.bold = True; lr.font.color.rgb = GREEN
    body = ''.join(frags)
    xml = f'<m:oMath {MNS}>{body}</m:oMath>'
    p._p.append(parse_xml(xml))
    tcPr = c._tc.get_or_add_tcPr(); borders = OxmlElement('w:tcBorders')
    left = OxmlElement('w:left'); left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '18'); left.set(qn('w:space'), '0'); left.set(qn('w:color'), '1E6B3A')
    borders.append(left); tcPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def answer_text(doc, text, label='Respuesta:'):
    t = doc.add_table(rows=1, cols=1); t.columns[0].width = Cm(16)
    c = t.cell(0, 0); _shade(c, 'E7F3EC'); c.width = Cm(16)
    p = c.paragraphs[0]
    lr = p.add_run(label + '  '); lr.bold = True; lr.font.color.rgb = GREEN
    vr = p.add_run(text); vr.bold = True
    tcPr = c._tc.get_or_add_tcPr(); borders = OxmlElement('w:tcBorders')
    left = OxmlElement('w:left'); left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '18'); left.set(qn('w:space'), '0'); left.set(qn('w:color'), '1E6B3A')
    borders.append(left); tcPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def step(doc, n, text):
    """A brief, numbered step explanation (concise) rendered above an equation."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Cm(0.3)
    lab = p.add_run(f"Step {n}.  "); lab.bold = True; lab.font.size = Pt(10.5)
    lab.font.color.rgb = ACCENT
    tr = p.add_run(text); tr.font.size = Pt(10.5); tr.font.color.rgb = RGBColor(0x33,0x33,0x33)
    return p

def image(doc, path, width_cm=10.5, caption=None):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Cm(width_cm))
    if caption:
        cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cp.add_run(caption); cr.italic = True; cr.font.size = Pt(9)
        cr.font.color.rgb = GREY
        cp.paragraph_format.space_after = Pt(6)
